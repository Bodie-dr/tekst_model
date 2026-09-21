from pathlib import Path
import hashlib
import re
from datetime import datetime, timezone

from docx import Document
from PIL import Image


def clean_text(text: str) -> str:
    text = text.replace(
        "\u00a0",
        " "
    )

    text = text.replace(
        "\r\n",
        "\n"
    )

    text = text.replace(
        "\r",
        "\n"
    )

    cleaned_lines = []

    for line in text.splitlines():
        line = re.sub(
            r"[ \t]+",
            " ",
            line
        ).strip()

        if line:
            cleaned_lines.append(line)

    return "\n".join(cleaned_lines).strip()


def read_docx(
    file_path: Path
) -> str:
    try:
        document = Document(file_path)

    except Exception as error:
        raise RuntimeError(
            "Het Word-document kan niet worden geopend: "
            f"{file_path.name}"
        ) from error

    text_parts = []

    # Paragrafen uitlezen

    for paragraph in document.paragraphs:
        paragraph_text = clean_text(
            paragraph.text
        )

        if paragraph_text:
            text_parts.append(
                paragraph_text
            )

    # Tabellen uitlezen

    for table_number, table in enumerate(
        document.tables,
        start=1
    ):
        table_rows = []

        for row in table.rows:
            cells = [
                clean_text(cell.text)
                for cell in row.cells
            ]

            if any(cells):
                table_rows.append(
                    " | ".join(cells)
                )

        if table_rows:
            text_parts.append(
                f"[Tabel {table_number}]\n"
                + "\n".join(table_rows)
            )

    return "\n\n".join(text_parts).strip()


def read_image_metadata(
    file_path: Path,
) -> str:
    try:
        with Image.open(file_path) as image:
            width, height = image.size
            image_format = image.format or file_path.suffix.lstrip(".").upper()

    except Exception as error:
        raise RuntimeError(
            "De afbeelding kan niet worden geopend: "
            f"{file_path.name}"
        ) from error

    return (
        "[Afbeelding]\n"
        f"Bestandsnaam: {file_path.name}\n"
        f"Formaat: {image_format}\n"
        f"Afmetingen: {width} x {height} pixels"
    )


def calculate_sha256(
    file_path: Path
) -> str:
    checksum = hashlib.sha256()

    with file_path.open("rb") as file_handle:
        while True:
            data = file_handle.read(
                1024 * 1024
            )

            if not data:
                break

            checksum.update(data)

    return checksum.hexdigest()


def get_modified_datetime(
    file_path: Path
) -> datetime:
    timestamp = file_path.stat().st_mtime

    return datetime.fromtimestamp(
        timestamp,
        tz=timezone.utc
    )